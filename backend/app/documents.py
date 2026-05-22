from hashlib import sha256
import mimetypes
from pathlib import Path
import re
from uuid import uuid4

from fastapi import UploadFile
from docx import Document as DocxDocument
from pypdf import PdfReader
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.audit import write_audit_log
from app.config import get_settings, resolve_config_path
from app.models import Document, SourceType
from app.text_normalization import normalize_text_block


TEXT_EXTENSIONS = {".txt", ".md", ".tex"}
DOCX_EXTENSIONS = {".docx"}
SUPPORTED_LOCAL_IMPORT_EXTENSIONS = TEXT_EXTENSIONS | DOCX_EXTENSIONS | {".pdf"}
EVAL_FILENAME_PATTERN = re.compile(r"^profile_sample_\d+\.(txt|md)$", re.IGNORECASE)


def _storage_root() -> Path:
    settings = get_settings()
    root = resolve_config_path(settings.local_storage_dir)
    root.mkdir(parents=True, exist_ok=True)
    return root


def _allowed_local_document_roots() -> list[Path]:
    settings = get_settings()
    roots: list[Path] = []
    for raw_root in settings.local_document_allowed_roots.split(","):
        candidate = raw_root.strip()
        if not candidate:
            continue
        roots.append(resolve_config_path(candidate))
    return roots


def validate_local_document_path(local_path: str) -> Path:
    path = Path(local_path).expanduser().resolve()
    if not path.exists():
        raise ValueError(f"Local document not found at {path}")
    if not path.is_file():
        raise ValueError(f"Local document path must point to a file: {path}")
    if path.suffix.lower() not in SUPPORTED_LOCAL_IMPORT_EXTENSIONS:
        supported = ", ".join(sorted(SUPPORTED_LOCAL_IMPORT_EXTENSIONS))
        raise ValueError(f"Unsupported local document type {path.suffix or '<none>'}. Allowed: {supported}")

    allowed_roots = _allowed_local_document_roots()
    if not allowed_roots:
        raise ValueError("Local document import is disabled because no allowed roots are configured.")

    for root in allowed_roots:
        try:
            path.relative_to(root)
            return path
        except ValueError:
            continue

    allowed_display = ", ".join(str(root) for root in allowed_roots)
    raise ValueError(
        "Local document import is restricted to approved folders. "
        f"Allowed roots: {allowed_display}"
    )


def is_evaluation_artifact_document(document: Document | None) -> bool:
    if document is None:
        return False

    metadata = document.document_metadata or {}
    if metadata.get("eval_artifact") is True:
        return True

    filename = (document.original_filename or "").strip()
    if EVAL_FILENAME_PATTERN.match(filename):
        return True

    storage_path = (document.storage_path or "").lower()
    if "tmp_eval_docs" in storage_path:
        return True

    return False


def _safe_filename(filename: str) -> str:
    return Path(filename).name.replace(" ", "_")


def _extract_pdf_text(path: Path) -> tuple[str, dict]:
    reader = PdfReader(str(path))
    pages: list[str] = []
    for index, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        pages.append(text.strip())

    text = "\n\n".join(page for page in pages if page)
    return normalize_text_block(text) or "", {"page_count": len(reader.pages)}


def _extract_text(path: Path, content_type: str | None) -> tuple[str | None, dict]:
    suffix = path.suffix.lower()
    if suffix == ".pdf" or content_type == "application/pdf":
        return _extract_pdf_text(path)
    if suffix in DOCX_EXTENSIONS or content_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        document = DocxDocument(str(path))
        paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
        return normalize_text_block("\n\n".join(paragraphs)), {"parser": "docx", "paragraph_count": len(paragraphs)}
    if suffix in TEXT_EXTENSIONS:
        return normalize_text_block(path.read_text(encoding="utf-8")), {"parser": "plain_text"}
    return None, {"parser": "unsupported"}


def _store_document_bytes(
    db: Session,
    *,
    filename: str,
    content: bytes,
    source_type: SourceType,
    content_type: str | None,
) -> Document:
    checksum = sha256(content).hexdigest()
    existing = db.scalar(select(Document).where(Document.checksum == checksum))
    if existing:
        return existing

    documents_dir = _storage_root() / "documents"
    documents_dir.mkdir(parents=True, exist_ok=True)

    storage_name = f"{uuid4()}_{_safe_filename(filename or 'document')}"
    storage_path = documents_dir / storage_name
    storage_path.write_bytes(content)

    extracted_text, metadata = _extract_text(storage_path, content_type)
    document = Document(
        source_type=source_type,
        original_filename=filename or storage_name,
        storage_path=str(storage_path),
        content_type=content_type,
        checksum=checksum,
        extracted_text=extracted_text,
        document_metadata={
            **metadata,
            "size_bytes": len(content),
            "original_content_type": content_type,
        },
    )
    db.add(document)
    db.flush()
    write_audit_log(
        db,
        event_type="document.uploaded",
        entity_type="document",
        entity_id=document.id,
        details={
            "source_type": source_type,
            "original_filename": document.original_filename,
            "checksum": checksum,
        },
    )
    db.commit()
    db.refresh(document)
    return document


async def store_uploaded_document(db: Session, *, file: UploadFile, source_type: SourceType) -> Document:
    content = await file.read()
    return _store_document_bytes(
        db,
        filename=file.filename or "document",
        content=content,
        source_type=source_type,
        content_type=file.content_type,
    )


def store_local_document(db: Session, *, source_type: SourceType, local_path: str) -> Document:
    path = validate_local_document_path(local_path)
    content_type, _ = mimetypes.guess_type(str(path))
    return _store_document_bytes(
        db,
        filename=path.name,
        content=path.read_bytes(),
        source_type=source_type,
        content_type=content_type,
    )


def list_documents(db: Session) -> list[Document]:
    return [
        document
        for document in db.scalars(select(Document).order_by(Document.created_at.desc()))
        if not is_evaluation_artifact_document(document)
    ]


def delete_document(db: Session, document_id) -> Document | None:
    document = db.get(Document, document_id)
    if document is None:
        return None
    if is_evaluation_artifact_document(document):
        return None

    storage_path = Path(document.storage_path).expanduser()
    if storage_path.exists() and storage_path.is_file():
        storage_path.unlink()

    db.delete(document)
    write_audit_log(
        db,
        event_type="document.deleted",
        entity_type="document",
        entity_id=document.id,
        details={
            "source_type": document.source_type.value if hasattr(document.source_type, "value") else str(document.source_type),
            "original_filename": document.original_filename,
        },
    )
    db.commit()
    return document
